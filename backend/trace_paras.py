import json, re, sys
import generate_docx_from_json as g
from docx import Document

with open("output/20260227_005756_Real_Time_LiDAR_SLAM_with_Reinforcement_Learning_P.json",encoding="utf-8") as f:
    paper = json.load(f)

intro = paper["sections"][0]["content"]
print(f"Content: {len(intro)} chars, bullets: {intro.count(chr(8226))}")

doc = Document()

# Patch add_paragraph to trace calls
import traceback as tb
call_count = [0]
orig = doc.add_paragraph

def patched_add_paragraph(text="", style=None):
    call_count[0] += 1
    n = call_count[0]
    if n <= 20 or n % 100 == 0:
        stack = tb.extract_stack()
        funcs = [f.name for f in stack[-4:-1]]
        print(f"  para#{n}: {funcs}")
    return orig(text, style)

doc.add_paragraph = patched_add_paragraph

g._add_content_paragraphs(doc, intro, ".", paper)
print(f"Total paras: {call_count[0]}")
